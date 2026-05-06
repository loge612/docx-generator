# -*- coding: utf-8 -*-
import os
import re
import uuid
import json
import base64
from typing import List, Optional, Union
from urllib.parse import urlparse, quote, unquote
from datetime import datetime, timedelta

import requests
from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tc, CT_Tbl
from docx.oxml.shared import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL

HOST = os.environ.get("HOST", "0.0.0.0")
PORT = int(os.environ.get("PORT", "8000"))
BASE_URL = os.environ.get("BASE_URL", f"http://localhost:{PORT}")

file_store = {}

app = FastAPI(
    title="制式文档生成API",
    description="Word模板占位符替换服务",
    version="1.2.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class GenerateRequest(BaseModel):
    template_file_url: str = "https://raw.githubusercontent.com/3293203411/docx-generator/main/gaoqilixiangmoban.docx"
    text_keys: Union[List[str], str]
    text_values: Union[List[str], str]
    filename: Optional[str] = None


def parse_json_param(param):
    if param is None:
        return []
    if isinstance(param, list):
        return [str(item) for item in param]
    if isinstance(param, str):
        try:
            parsed = json.loads(param)
            if isinstance(parsed, list):
                return [str(item) for item in parsed]
            return [str(parsed)]
        except json.JSONDecodeError:
            return [param]
    return [str(param)]


def download_file(url: str) -> bytes:
    try:
        headers = {'User-Agent': 'Mozilla/5.0'}
        response = requests.get(url, headers=headers, timeout=60)
        response.raise_for_status()
        return response.content
    except requests.exceptions.RequestException as e:
        raise HTTPException(status_code=400, detail=f"下载模板文件失败: {str(e)}")


def merge_runs_in_paragraph(paragraph):
    if not paragraph.runs:
        return {"text": "", "runs": [], "formatting": None}
    full_text = ""
    for run in paragraph.runs:
        full_text += run.text if run.text else ""
    first_run = paragraph.runs[0]
    formatting = {
        "bold": first_run.bold,
        "italic": first_run.italic,
        "underline": first_run.underline,
        "font.name": first_run.font.name,
        "font.size": first_run.font.size,
        "font.color.rgb": first_run.font.color.rgb if first_run.font.color and first_run.font.color.rgb else None,
    }
    return {"text": full_text, "runs": list(paragraph.runs), "formatting": formatting}


def replace_placeholders_in_text(text, keys, values):
    result = text
    for key, value in zip(keys, values):
        pattern = r'\{\{\s*' + re.escape(key) + r'\s*\}\}'
        result = re.sub(pattern, str(value), result)
    return result


def apply_formatting_to_run(run, formatting):
    if formatting is None:
        return
    if formatting.get("bold") is not None:
        run.bold = formatting["bold"]
    if formatting.get("italic") is not None:
        run.italic = formatting["italic"]
    if formatting.get("underline") is not None:
        run.underline = formatting["underline"]
    if formatting.get("font.name"):
        run.font.name = formatting["font.name"]
    if formatting.get("font.size"):
        run.font.size = formatting["font.size"]
    if formatting.get("font.color.rgb"):
        run.font.color.rgb = formatting["font.color.rgb"]


def parse_colored_segments(value):
    """
    解析 <red>文字</red> 标记
    返回 [(文字, is_red), ...]
    示例：
        "hello <red>世界</red> end"
        → [("hello ", False), ("世界", True), (" end", False)]
    """
    segments = []
    pattern = re.compile(r'<red>(.*?)</red>', re.DOTALL)
    last_end = 0

    for match in pattern.finditer(value):
        if match.start() > last_end:
            segments.append((value[last_end:match.start()], False))
        segments.append((match.group(1), True))
        last_end = match.end()

    if last_end < len(value):
        segments.append((value[last_end:], False))

    if not segments:
        segments = [(value, False)]

    return segments


def has_color_tag(text):
    """判断文本中是否含有<red>标签"""
    return '<red>' in text and '</red>' in text


def set_cell_border(cell):
    """设置单元格边框"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tcBorders.append(border)
    tcPr.append(tcBorders)


def set_cell_background(cell, color):
    """设置单元格背景色"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, font_size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER):
    """设置单元格文字"""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    para = cell.paragraphs[0]
    para.alignment = align
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def build_research_table(doc, table_data):
    """
    构建研发内容完成情况表格
    table_data 格式:
    [
        ['内容1', '高吸附容量分子筛材料研发', '已完成', '100%', '—'],
        ['内容2', '低能耗催化燃烧工艺优化',   '已完成', '100%', '—'],
        ...
    ]
    """
    headers = ['序号', '原计划研发内容', '实际完成情况', '完成度', '未完成原因及说明']

    table = doc.add_table(rows=1 + len(table_data), cols=len(headers))
    table.style = 'Table Grid'

    col_widths = [Inches(0.8), Inches(2.5), Inches(1.5), Inches(1.0), Inches(2.0)]
    for col_idx, width in enumerate(col_widths):
        for row in table.rows:
            row.cells[col_idx].width = width

    for col_idx, header in enumerate(headers):
        cell = table.cell(0, col_idx)
        set_cell_background(cell, 'D9E1F2')
        set_cell_border(cell)
        set_cell_text(cell, header, bold=True)

    for row_idx, row_data in enumerate(table_data):
        for col_idx, value in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            set_cell_border(cell)
            if row_idx % 2 == 1:
                set_cell_background(cell, 'F2F2F2')
            set_cell_text(cell, value)

    return table


def replace_placeholder_with_table(doc, placeholder, table_data):
    """
    找到含占位符的段落，替换为真实表格
    """
    target_para = None
    target_idx = None

    body = doc.element.body
    body_elements = list(body)

    for idx, element in enumerate(body_elements):
        if isinstance(element, CT_P):
            para = Paragraph(element, doc)
            full_text = "".join(run.text for run in para.runs)
            if placeholder in full_text:
                target_para = para
                target_idx = idx
                break

    if target_para is None:
        return False

    tmp_doc = Document()
    tmp_table = build_research_table(tmp_doc, table_data)
    tbl_element = tmp_table._tbl

    target_para._element.addprevious(tbl_element)

    parent = target_para._element.getparent()
    parent.remove(target_para._element)

    return True


def process_paragraph(paragraph, keys, values):
    if not paragraph.runs:
        return
    para_info = merge_runs_in_paragraph(paragraph)
    original_text = para_info["text"]
    if not original_text:
        return

    # 检查是否有占位符
    has_placeholder = False
    for key in keys:
        if re.search(r'\{\{\s*' + re.escape(key) + r'\s*\}\}', original_text):
            has_placeholder = True
            break
    if not has_placeholder:
        return

    # 替换后的文本
    new_text = replace_placeholders_in_text(original_text, keys, values)
    if new_text == original_text:
        return

    # ============================================
    # 判断是否需要处理红色标签
    # ============================================
    if not has_color_tag(new_text):
        # ✅ 原来的逻辑：没有红色标签，直接替换
        for run in paragraph.runs:
            run.text = ""
        lines = new_text.split('\n')
        current_run = paragraph.runs[0]
        for i, line in enumerate(lines):
            if i == 0:
                current_run.text = line
            else:
                current_run.add_break()
                current_run.text += line
        apply_formatting_to_run(current_run, para_info["formatting"])

    else:
        # ✅ 新逻辑：含有<red>标签，拆分成多个run
        formatting = para_info["formatting"]

        # 清空所有现有runs
        p_elem = paragraph._p
        for run in paragraph.runs:
            p_elem.remove(run._r)

        # 按换行拆分
        lines = new_text.split('\n')

        for line_idx, line in enumerate(lines):
            # 每行解析红色片段
            segments = parse_colored_segments(line)

            for seg_text, is_red in segments:
                if not seg_text:
                    continue

                # 创建新run
                new_run = OxmlElement('w:r')

                # 复制格式 rPr
                rPr = OxmlElement('w:rPr')

                # 字体
                if formatting.get("font.name"):
                    rFonts = OxmlElement('w:rFonts')
                    rFonts.set(qn('w:ascii'), formatting["font.name"])
                    rFonts.set(qn('w:hAnsi'), formatting["font.name"])
                    rFonts.set(qn('w:eastAsia'), formatting["font.name"])
                    rPr.append(rFonts)

                # 加粗
                if formatting.get("bold"):
                    bold_elem = OxmlElement('w:b')
                    rPr.append(bold_elem)

                # 斜体
                if formatting.get("italic"):
                    italic_elem = OxmlElement('w:i')
                    rPr.append(italic_elem)

                # 下划线
                if formatting.get("underline"):
                    u_elem = OxmlElement('w:u')
                    u_elem.set(qn('w:val'), 'single')
                    rPr.append(u_elem)

                # 字号
                if formatting.get("font.size"):
                    try:
                        font_size = formatting["font.size"]
                        sz_val = str(int(font_size.pt * 2))
                        sz = OxmlElement('w:sz')
                        sz.set(qn('w:val'), sz_val)
                        rPr.append(sz)
                        szCs = OxmlElement('w:szCs')
                        szCs.set(qn('w:val'), sz_val)
                        rPr.append(szCs)
                    except Exception:
                        pass

                # 颜色：红色 or 原色
                color_elem = OxmlElement('w:color')
                if is_red:
                    color_elem.set(qn('w:val'), 'FF0000')
                elif formatting.get("font.color.rgb"):
                    color_elem.set(qn('w:val'), str(formatting["font.color.rgb"]))
                else:
                    color_elem.set(qn('w:val'), 'auto')
                rPr.append(color_elem)

                new_run.append(rPr)

                # 文字内容
                t_elem = OxmlElement('w:t')
                t_elem.text = seg_text
                # 保留首尾空格
                if seg_text.startswith(' ') or seg_text.endswith(' '):
                    t_elem.set(
                        '{http://www.w3.org/XML/1998/namespace}space',
                        'preserve'
                    )
                new_run.append(t_elem)
                p_elem.append(new_run)

            # 换行处理（非最后一行）
            if line_idx < len(lines) - 1:
                br_run = OxmlElement('w:r')
                br = OxmlElement('w:br')
                br_run.append(br)
                p_elem.append(br_run)


def process_table(table, keys, values):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                process_paragraph(paragraph, keys, values)


def process_document(doc, keys, values):
    """
    处理文档：
    1. 先处理表格类占位符（替换为真实表格）
    2. 再处理普通文本占位符
    """
    # =============================
    # ✅ 第一步：处理需要转表格的占位符
    # =============================
    table_placeholder = '研发内容完成情况'
    table_data_key_idx = None
    table_data = None

    for i, key in enumerate(keys):
        if key == table_placeholder:
            table_data_key_idx = i
            break

    if table_data_key_idx is not None:
        raw_value = values[table_data_key_idx]
        try:
            table_data = json.loads(raw_value)
        except Exception:
            table_data = None

        if table_data:
            replace_placeholder_with_table(doc, f'{{{{{table_placeholder}}}}}', table_data)
            keys = [k for i, k in enumerate(keys) if i != table_data_key_idx]
            values = [v for i, v in enumerate(values) if i != table_data_key_idx]

    # =============================
    # ✅ 第二步：处理普通文本占位符
    # =============================
    for element in doc.element.body:
        if isinstance(element, CT_P):
            paragraph = Paragraph(element, doc)
            process_paragraph(paragraph, keys, values)

    for table in doc.tables:
        process_table(table, keys, values)


def generate_filename(original_url, custom_filename=None):
    if custom_filename:
        if not custom_filename.lower().endswith('.docx'):
            custom_filename += '.docx'
        return custom_filename
    parsed = urlparse(original_url)
    basename = os.path.basename(parsed.path)
    if basename and basename.lower().endswith('.docx'):
        unique_id = uuid.uuid4().hex[:8]
        name, ext = os.path.splitext(basename)
        return f"{name}_{unique_id}{ext}"
    return f"generated_{uuid.uuid4().hex[:8]}.docx"


@app.get("/")
async def root():
    return {"service": "制式文档生成API", "version": "1.2.0", "status": "running"}


@app.get("/health")
async def health_check():
    return {"status": "healthy"}


@app.post("/generate")
async def generate_document(request: GenerateRequest):
    try:
        keys = parse_json_param(request.text_keys)
        values = parse_json_param(request.text_values)
        if len(keys) != len(values):
            raise HTTPException(
                status_code=400,
                detail=f"text_keys数量({len(keys)})与text_values数量({len(values)})不匹配"
            )
        if not keys:
            raise HTTPException(status_code=400, detail="text_keys不能为空")

        file_content = download_file(request.template_file_url)

        import tempfile
        temp_input = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
        temp_input.write(file_content)
        temp_input.close()

        try:
            doc = Document(temp_input.name)
            process_document(doc, keys, values)

            output_filename = generate_filename(request.template_file_url, request.filename)
            encoded_filename = quote(output_filename, encoding="utf-8")

            import io
            buffer = io.BytesIO()
            doc.save(buffer)
            file_bytes = buffer.getvalue()

            file_store[encoded_filename] = {
                "data": file_bytes,
                "expires": datetime.now() + timedelta(hours=2),
                "real_name": output_filename
            }

            download_url = f"{BASE_URL}/download/{encoded_filename}"

            return JSONResponse({
                "success": True,
                "message": "文档生成成功",
                "full_download_url": download_url,
                "filename": output_filename,
                "replaced_count": len(keys)
            })
        finally:
            os.unlink(temp_input.name)

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"生成文档失败: {str(e)}")


@app.get("/download/{filename}")
async def download_file_endpoint(filename: str):
    filename = os.path.basename(filename)

    if filename in file_store:
        item = file_store[filename]
        if datetime.now() < item["expires"]:
            real_name = item["real_name"]
            content_disposition = f"attachment; filename*=utf-8''{quote(real_name)}"
            return Response(
                content=item["data"],
                media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                headers={"Content-Disposition": content_disposition}
            )
        else:
            del file_store[filename]
            raise HTTPException(status_code=404, detail="文件已过期")

    raise HTTPException(status_code=404, detail="文件不存在")


def start_server():
    import uvicorn
    uvicorn.run(app, host=HOST, port=PORT)


if __name__ == "__main__":
    start_server()
